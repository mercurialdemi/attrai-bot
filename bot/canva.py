"""File modificabile per Canva (PowerPoint): Canva lo importa mantenendo i testi modificabili.

- Slide di solo testo: sfondo pulito come immagine + una casella di testo per ogni paragrafo/blocco (le righe
  dello stesso blocco stanno nella stessa casella, con gli stessi a capo, posizione, dimensione e colore).
  In Canva si clicca sul blocco e si modifica tutta la frase insieme.
- Slide con grafica (rifatte da Gemini) o rimaste uguali all'originale: l'immagine finale.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

PX = 9525          # 1 pixel a 96 dpi in EMU
FONT_NAMES = {"regular": "Inter", "semibold": "Inter SemiBold", "bold": "Inter", "titolo": "Anton", "stretto": "Oswald",
              "regular_tight": "Inter Tight", "semibold_tight": "Inter Tight SemiBold", "bold_tight": "Inter Tight",
              "regular_montserrat": "Montserrat", "semibold_montserrat": "Montserrat SemiBold", "bold_montserrat": "Montserrat",
              "regular_poppins": "Poppins", "semibold_poppins": "Poppins SemiBold", "bold_poppins": "Poppins",
              "regular_arimo": "Arimo", "semibold_arimo": "Arimo SemiBold", "bold_arimo": "Arimo",
              "regular_gaegu": "Gaegu", "semibold_gaegu": "Gaegu", "bold_gaegu": "Gaegu",
              "regular_schoolbell": "Schoolbell", "semibold_schoolbell": "Schoolbell", "bold_schoolbell": "Schoolbell",
              "regular_architects": "Architects Daughter", "semibold_architects": "Architects Daughter", "bold_architects": "Architects Daughter"}


def build(post: dict, paths: list[Path], out_path: Path, failed: list[int] | None = None) -> Path:
    from PIL import Image

    failed = set(failed or [])
    first = Image.open(paths[0])
    W, H = first.size
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
    blank = prs.slide_layouts[6]

    for i, p in enumerate(paths, 1):
        if i in failed:
            continue
        slide_info = post["slides"][i - 1] if i - 1 < len(post.get("slides", [])) else {}
        canva = slide_info.get("_canva") or {}
        s = prs.slides.add_slide(blank)
        if canva.get("tipo") == "testo" and canva.get("sfondo") and Path(canva["sfondo"]).exists():
            s.shapes.add_picture(canva["sfondo"], 0, 0, Emu(W * PX), Emu(H * PX))
            for g in _blocchi(canva.get("righe", []), W):
                _casella(s, g, W)
        else:
            s.shapes.add_picture(str(p), 0, 0, Emu(W * PX), Emu(H * PX))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def _larghezza(r: dict) -> float:
    try:
        from .textslide import _font
        return float(_font(r["weight"], int(r["size"])).getlength(r["testo"]))
    except Exception:
        return len(r["testo"]) * r["size"] * 0.5


def _blocchi(righe: list[dict], W: int) -> list[dict]:
    """Raggruppa le righe dello stesso paragrafo (stesso corpo e peso, una sotto l'altra a passo regolare,
    stesso bordo sinistro oppure stesso centro) in un'unica casella di testo."""
    gruppi: list[dict] = []
    for r in righe:
        if not (r.get("testo") or "").strip():
            continue
        w = _larghezza(r)
        c = r["x"] + w / 2
        g = gruppi[-1] if gruppi else None
        if g:
            last = g["righe"][-1]
            passo = r["y"] - last["y"]
            stesso_stile = r["size"] == last["size"] and r["weight"] == last["weight"]
            vicina = 0 < passo <= r["size"] * 1.9 and (g["passo"] is None or abs(passo - g["passo"]) <= r["size"] * 0.25)
            allineata = (abs(r["x"] - g["x"]) <= 4 and g["allin"] in (None, "sinistra")) or \
                        (abs(c - g["centro"]) <= 6 and g["allin"] in (None, "centro"))
            if stesso_stile and vicina and allineata:
                if g["allin"] is None:
                    g["allin"] = "sinistra" if abs(r["x"] - g["x"]) <= 4 else "centro"
                g["passo"] = g["passo"] or passo
                g["righe"].append(r)
                g["x"] = min(g["x"], r["x"])
                g["destra"] = max(g["destra"], r["x"] + w)
                continue
        gruppi.append({"righe": [r], "x": r["x"], "destra": r["x"] + w, "centro": c, "passo": None, "allin": None})
    for g in gruppi:
        if g["allin"] is None:   # riga sola: centrata se sta nel mezzo della slide
            g["allin"] = "centro" if abs(g["centro"] - W / 2) <= W * 0.02 else "sinistra"
    return gruppi


def _casella(s, g: dict, W: int) -> None:
    righe = g["righe"]
    size = righe[0]["size"]
    passo = g["passo"] or size * 1.25
    margine = size * 0.6                       # un filo di spazio: in Canva la riga non va a capo da sola
    if g["allin"] == "centro":
        larg = (g["destra"] - g["x"]) + 2 * margine
        x0 = g["centro"] - larg / 2
    else:
        x0 = g["x"]
        larg = (g["destra"] - g["x"]) + margine
    x0 = max(0.0, x0)
    larg = min(larg, W - x0)
    alto = passo * (len(righe) - 1) + size * 1.3
    box = s.shapes.add_textbox(Emu(int(x0 * PX)), Emu(int(righe[0]["y"] * PX)), Emu(int(larg * PX)), Emu(int(alto * PX)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, r in enumerate(righe):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = PP_ALIGN.CENTER if g["allin"] == "centro" else PP_ALIGN.LEFT
        par.line_spacing = Pt(passo * 0.75)    # stessa distanza tra le righe della slide
        runs = r.get("runs") or [[r["testo"], r["color"]]]
        for k, item in enumerate(runs):
            word, col = item[0], item[1]
            strong = bool(item[2]) if len(item) > 2 else False
            run = par.add_run()
            run.text = word + (" " if k < len(runs) - 1 else "")
            f = run.font
            f.name = FONT_NAMES.get(r["weight"], "Inter")
            f.size = Pt(r["size"] * 0.75)          # pixel -> punti
            f.bold = r["weight"].startswith("bold") or strong
            f.color.rgb = RGBColor(*list(col)[:3])
